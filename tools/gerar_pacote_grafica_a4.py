#!/usr/bin/env python3
"""Cria um pacote A4 a 300 dpi para graficas e servicos de fotolivros."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
import textwrap
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageCms, ImageDraw, ImageFilter, ImageFont, ImageOps


ROOT = Path(__file__).resolve().parents[1]
BOOKS_DIR = ROOT / "books"
PRINT_DIR = ROOT / "impressao"
DELIVERY_DIR = ROOT / "entregas"
PACKAGE_DIR = DELIVERY_DIR / "pacote-impressao-a4"
PACKAGE_ZIP = DELIVERY_DIR / "as-historias-do-cao-joaquim-pacote-grafica-a4.zip"
COLLECTION_PDF = PRINT_DIR / "as-historias-do-cao-joaquim.pdf"
COVER_ART = ROOT / "assets" / "capa-coletanea-original.png"

A4_PORTRAIT = (2480, 3508)
A4_LANDSCAPE = (3508, 2480)
DPI = (300, 300)
BLEED_PX = round(3 / 25.4 * 300)
SPINE_WIDTHS_MM = (10, 15, 20, 25, 30)

PAPER = "#f3eadb"
INK = "#1d1b19"
MUTED = "#59605b"
LINE = "#aaa294"
BRICK = "#b54733"
PETROL = "#315f63"
BLUE = "#16719a"

WINDOWS_FONTS = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
FONT_REGULAR = WINDOWS_FONTS / "georgia.ttf"
FONT_BOLD = WINDOWS_FONTS / "georgiab.ttf"
FONT_ITALIC = WINDOWS_FONTS / "georgiai.ttf"

SRGB_PROFILE = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()

SITE_ROOT_FILES = ("index.html", "reader.html", "direitos.html", "LICENSE")
SITE_DIRECTORIES = ("assets", "books", "v1", "v2")


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    if not path.exists():
        raise FileNotFoundError(f"Tipo de letra em falta: {path}")
    return ImageFont.truetype(str(path), size=size)


def save_jpeg(image: Image.Image, path: Path, quality: int = 96) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.convert("RGB").save(
        path,
        "JPEG",
        quality=quality,
        subsampling=0,
        optimize=True,
        progressive=True,
        dpi=DPI,
        icc_profile=SRGB_PROFILE,
    )


def save_png(image: Image.Image, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, "PNG", optimize=True, dpi=DPI, icc_profile=SRGB_PROFILE)


def fit_on_canvas(
    source: Image.Image,
    size: tuple[int, int],
    margin: int,
    background: str = PAPER,
    border: bool = True,
) -> Image.Image:
    canvas = Image.new("RGB", size, background)
    available = (size[0] - margin * 2, size[1] - margin * 2)
    fitted = ImageOps.contain(source.convert("RGB"), available, Image.Resampling.LANCZOS)
    if fitted.width > source.width or fitted.height > source.height:
        fitted = fitted.filter(ImageFilter.UnsharpMask(radius=1.2, percent=105, threshold=3))
    x = (size[0] - fitted.width) // 2
    y = (size[1] - fitted.height) // 2
    if border:
        draw = ImageDraw.Draw(canvas)
        draw.rectangle((x - 18, y - 18, x + fitted.width + 17, y + fitted.height + 17), fill="#fffaf0")
        draw.rectangle((x - 20, y - 20, x + fitted.width + 19, y + fitted.height + 19), outline=LINE, width=3)
    canvas.paste(fitted, (x, y))
    return canvas


def wrap_line(draw: ImageDraw.ImageDraw, text: str, typeface: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if draw.textlength(candidate, font=typeface) <= width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def text_layout(
    draw: ImageDraw.ImageDraw,
    paragraphs: list[str],
    max_width: int,
    max_height: int,
) -> tuple[ImageFont.FreeTypeFont, list[list[str]], int, int]:
    for size in range(64, 41, -2):
        typeface = font(FONT_REGULAR, size)
        lines = [wrap_line(draw, paragraph, typeface, max_width) for paragraph in paragraphs]
        line_height = round(size * 1.48)
        paragraph_gap = round(size * 0.62)
        total_height = sum(len(block) * line_height + paragraph_gap for block in lines) - paragraph_gap
        if total_height <= max_height:
            return typeface, lines, line_height, paragraph_gap
    raise ValueError("O texto nao cabe numa pagina A4.")


def draw_text_image(book: dict, page: dict, display_page: int, output: Path) -> None:
    canvas = Image.new("RGB", A4_PORTRAIT, PAPER)
    draw = ImageDraw.Draw(canvas)
    inset = 150
    draw.rectangle((inset, inset, A4_PORTRAIT[0] - inset, A4_PORTRAIT[1] - inset), outline=LINE, width=3)

    triangles = (
        [(150, 150), (340, 150), (150, 340)],
        [(A4_PORTRAIT[0] - 150, A4_PORTRAIT[1] - 150),
         (A4_PORTRAIT[0] - 340, A4_PORTRAIT[1] - 150),
         (A4_PORTRAIT[0] - 150, A4_PORTRAIT[1] - 340)],
    )
    draw.polygon(triangles[0], fill=BRICK)
    draw.polygon(triangles[1], fill=PETROL)
    draw.polygon([(150, 150), (250, 150), (150, 250)], fill=INK)

    paragraphs = list(page.get("text", []))
    if page.get("end"):
        paragraphs.append(page["end"])
    typeface, wrapped, line_height, paragraph_gap = text_layout(draw, paragraphs, 1740, 2250)
    total_height = sum(len(block) * line_height + paragraph_gap for block in wrapped) - paragraph_gap
    y = max(560, (A4_PORTRAIT[1] - total_height) // 2)
    x = 370

    for paragraph_index, block in enumerate(wrapped):
        for line in block:
            is_end = page.get("end") and paragraph_index == len(wrapped) - 1
            line_font = font(FONT_ITALIC, typeface.size) if is_end else typeface
            color = BRICK if is_end else INK
            if is_end:
                line_width = draw.textlength(line, font=line_font)
                draw.text(((A4_PORTRAIT[0] - line_width) / 2, y), line, font=line_font, fill=color)
            else:
                draw.text((x, y), line, font=line_font, fill=color)
            y += line_height
        y += paragraph_gap

    footer_font = font(FONT_REGULAR, 29)
    footer = f"{book['title']}   {display_page}"
    footer_width = draw.textlength(footer, font=footer_font)
    draw.text(((A4_PORTRAIT[0] - footer_width) / 2, A4_PORTRAIT[1] - 235), footer, font=footer_font, fill=MUTED)
    save_png(canvas, output)


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    box: tuple[int, int, int, int],
    typeface: ImageFont.FreeTypeFont,
    fill: str,
    spacing: int = 14,
) -> None:
    x1, y1, x2, y2 = box
    lines = wrap_line(draw, text, typeface, x2 - x1)
    heights = [draw.textbbox((0, 0), line, font=typeface)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        width = draw.textlength(line, font=typeface)
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=typeface, fill=fill)
        y += height + spacing


def make_front_cover(size: tuple[int, int]) -> Image.Image:
    art = Image.open(COVER_ART).convert("RGB")
    cover = ImageOps.fit(art, size, Image.Resampling.LANCZOS, centering=(0.5, 0.48))
    draw = ImageDraw.Draw(cover)
    panel = (round(size[0] * 0.39), round(size[1] * 0.055), round(size[0] * 0.94), round(size[1] * 0.28))
    draw.rectangle(panel, fill="#f6eddf")
    draw.rectangle(panel, outline="#d9cebd", width=max(3, size[0] // 700))
    draw_centered_text(
        draw,
        "As Histórias do Cão Joaquim",
        (panel[0] + 70, panel[1] + 35, panel[2] - 70, panel[3] - 110),
        font(FONT_BOLD, round(size[0] * 0.055)),
        INK,
        spacing=16,
    )
    subtitle = "Uma coleção de aventuras para ler em família"
    subtitle_font = font(FONT_ITALIC, round(size[0] * 0.021))
    subtitle_width = draw.textlength(subtitle, font=subtitle_font)
    draw.text(((panel[0] + panel[2] - subtitle_width) / 2, panel[3] - 100), subtitle, font=subtitle_font, fill=BRICK)

    author_font = font(FONT_BOLD, round(size[0] * 0.021))
    author = "Miguel Pinto"
    author_width = draw.textlength(author, font=author_font)
    author_y = size[1] - round(size[1] * 0.055)
    draw.rectangle((size[0] - author_width - 170, author_y - 22, size[0] - 60, author_y + author_font.size + 22), fill="#f6eddf")
    draw.text((size[0] - author_width - 112, author_y), author, font=author_font, fill=INK)
    return cover


def paper_texture(image: Image.Image) -> Image.Image:
    texture = Image.effect_noise(image.size, 7).convert("L")
    tint = Image.new("RGB", image.size, PAPER)
    return Image.blend(image.convert("RGB"), Image.composite(tint, image, texture), 0.08)


def make_back_cover(size: tuple[int, int], book_count: int) -> Image.Image:
    cover = Image.new("RGB", size, PETROL)
    draw = ImageDraw.Draw(cover)
    unit = max(40, size[0] // 18)
    draw.polygon([(0, 0), (unit * 5, 0), (0, unit * 5)], fill=INK)
    draw.polygon([(size[0], size[1]), (size[0] - unit * 6, size[1]), (size[0], size[1] - unit * 6)], fill=BRICK)
    draw.polygon([(0, size[1]), (unit * 4, size[1]), (0, size[1] - unit * 4)], fill=BLUE)

    circle = round(size[0] * 0.115)
    center = (size[0] // 2, round(size[1] * 0.25))
    draw.ellipse((center[0] - circle, center[1] - circle, center[0] + circle, center[1] + circle), fill=BRICK)
    mark_font = font(FONT_BOLD, round(size[0] * 0.055))
    mark = "CJ"
    mark_width = draw.textlength(mark, font=mark_font)
    mark_box = draw.textbbox((0, 0), mark, font=mark_font)
    draw.text((center[0] - mark_width / 2, center[1] - (mark_box[3] - mark_box[1]) / 2 - mark_box[1]), mark, font=mark_font, fill="#fffaf0")

    draw_centered_text(
        draw,
        "As Histórias do Cão Joaquim",
        (round(size[0] * 0.15), round(size[1] * 0.39), round(size[0] * 0.85), round(size[1] * 0.51)),
        font(FONT_BOLD, round(size[0] * 0.052)),
        "#fffaf0",
    )
    draw_centered_text(
        draw,
        f"{book_count} aventuras inventadas em família, cheias de imaginação, gargalhadas e descobertas.",
        (round(size[0] * 0.19), round(size[1] * 0.53), round(size[0] * 0.81), round(size[1] * 0.66)),
        font(FONT_REGULAR, round(size[0] * 0.027)),
        "#fffaf0",
        spacing=18,
    )
    draw_centered_text(
        draw,
        "Criado por Miguel Pinto para o Guilherme e o Xavier",
        (round(size[0] * 0.18), round(size[1] * 0.72), round(size[0] * 0.82), round(size[1] * 0.79)),
        font(FONT_ITALIC, round(size[0] * 0.022)),
        "#fffaf0",
    )
    copyright_font = font(FONT_REGULAR, round(size[0] * 0.014))
    copyright_text = "Texto, personagens e edição © 2026 Miguel Pinto. Todos os direitos reservados."
    copyright_width = draw.textlength(copyright_text, font=copyright_font)
    draw.text(((size[0] - copyright_width) / 2, size[1] - round(size[1] * 0.07)), copyright_text, font=copyright_font, fill="#fffaf0")
    return paper_texture(cover)


def make_spine(width_px: int, height_px: int) -> Image.Image:
    spine = Image.new("RGB", (width_px, height_px), BRICK)
    draw = ImageDraw.Draw(spine)
    draw.rectangle((0, 0, max(3, width_px // 18), height_px), fill=INK)
    draw.rectangle((width_px - max(3, width_px // 18), 0, width_px, height_px), fill=PETROL)

    title_font = font(FONT_BOLD, min(58, max(32, width_px // 3)))
    author_font = font(FONT_REGULAR, min(38, max(24, width_px // 5)))
    title_layer = Image.new("RGBA", (height_px - 500, width_px), (0, 0, 0, 0))
    title_draw = ImageDraw.Draw(title_layer)
    title = "As Histórias do Cão Joaquim"
    title_width = title_draw.textlength(title, font=title_font)
    title_draw.text(((title_layer.width - title_width) / 2, max(5, (width_px - title_font.size) / 2 - 5)), title, font=title_font, fill="#fffaf0")
    rotated = title_layer.rotate(90, expand=True)
    spine.alpha_composite(rotated.convert("RGBA"), ((width_px - rotated.width) // 2, 250)) if spine.mode == "RGBA" else spine.paste(rotated, ((width_px - rotated.width) // 2, 250), rotated)

    author_layer = Image.new("RGBA", (700, width_px), (0, 0, 0, 0))
    author_draw = ImageDraw.Draw(author_layer)
    author = "Miguel Pinto"
    author_width = author_draw.textlength(author, font=author_font)
    author_draw.text(((author_layer.width - author_width) / 2, max(5, (width_px - author_font.size) / 2 - 5)), author, font=author_font, fill="#fffaf0")
    author_rotated = author_layer.rotate(90, expand=True)
    spine.paste(author_rotated, ((width_px - author_rotated.width) // 2, height_px - author_rotated.height - 220), author_rotated)
    return spine.convert("RGB")


def make_wrap(front: Image.Image, back: Image.Image, spine_mm: int) -> Image.Image:
    spine_px = round(spine_mm / 25.4 * 300)
    width = BLEED_PX + A4_PORTRAIT[0] + spine_px + A4_PORTRAIT[0] + BLEED_PX
    height = A4_PORTRAIT[1] + BLEED_PX * 2
    wrap = Image.new("RGB", (width, height), PETROL)
    back_bleed = ImageOps.fit(back, (A4_PORTRAIT[0] + BLEED_PX, height), Image.Resampling.LANCZOS, centering=(0.52, 0.5))
    front_bleed = ImageOps.fit(front, (A4_PORTRAIT[0] + BLEED_PX, height), Image.Resampling.LANCZOS, centering=(0.48, 0.5))
    wrap.paste(back_bleed, (0, 0))
    spine = make_spine(spine_px, A4_PORTRAIT[1])
    wrap.paste(spine, (BLEED_PX + A4_PORTRAIT[0], BLEED_PX))
    wrap.paste(front_bleed, (BLEED_PX + A4_PORTRAIT[0] + spine_px, 0))
    return wrap


def render_collection_pages(output_dir: Path) -> list[Path]:
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("pdftoppm nao foi encontrado. Instala o Poppler para criar as paginas A4.")
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="cao-joaquim-a4-", dir=DELIVERY_DIR) as temp_dir:
        prefix = Path(temp_dir) / "pagina"
        subprocess.run(
            [
                pdftoppm,
                "-jpeg",
                "-scale-to-x",
                str(A4_PORTRAIT[0]),
                "-scale-to-y",
                str(A4_PORTRAIT[1]),
                "-jpegopt",
                "quality=100",
                str(COLLECTION_PDF),
                str(prefix),
            ],
            check=True,
        )
        rendered = sorted(Path(temp_dir).glob("pagina-*.jpg"))
        outputs: list[Path] = []
        for index, source in enumerate(rendered, start=1):
            target = output_dir / f"pagina-{index:03d}.jpg"
            image = Image.open(source).convert("RGB")
            if image.size != A4_PORTRAIT:
                image = image.resize(A4_PORTRAIT, Image.Resampling.LANCZOS)
            save_jpeg(image, target)
            outputs.append(target)
    return outputs


def generate_illustrations(books: list[tuple[Path, dict]], output_dir: Path) -> int:
    count = 0
    for book_index, (book_dir, book) in enumerate(books, start=1):
        story_dir = output_dir / f"{book_index:02d}-{book['id']}"
        image_number = 0
        for page_index, page in enumerate(book["pages"], start=1):
            if "image" not in page:
                continue
            image_number += 1
            source = Image.open(book_dir / page["image"]).convert("RGB")
            if page["type"] == "cover":
                canvas = fit_on_canvas(source, A4_PORTRAIT, 120)
                label = "capa"
            else:
                canvas = fit_on_canvas(source, A4_LANDSCAPE, 120)
                label = f"ilustracao-{image_number - 1:02d}"
            save_jpeg(canvas, story_dir / f"{page_index:02d}-{label}.jpg")
            count += 1
    return count


def generate_text_images(books: list[tuple[Path, dict]], output_dir: Path) -> int:
    count = 0
    for book_index, (_, book) in enumerate(books, start=1):
        story_dir = output_dir / f"{book_index:02d}-{book['id']}"
        for page_index, page in enumerate(book["pages"], start=1):
            if page["type"] not in {"text", "ending"}:
                continue
            draw_text_image(book, page, page_index, story_dir / f"pagina-{page_index:02d}-texto.png")
            count += 1
    return count


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def create_text_document(books: list[tuple[Path, dict]], output_path: Path) -> int:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(24)
    section.bottom_margin = Mm(24)
    section.left_margin = Mm(27)
    section.right_margin = Mm(27)

    styles = document.styles
    styles["Normal"].font.name = "Georgia"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    styles["Normal"].font.size = Pt(12.5)
    styles["Title"].font.name = "Georgia"
    styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.size = Pt(28)
    title_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    styles["Heading 1"].font.name = "Georgia"
    styles["Heading 1"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.name = "Georgia"
    styles["Heading 2"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    styles["Heading 2"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 2"].font.size = Pt(12)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("As Histórias do Cão Joaquim")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(20)
    run = subtitle.add_run("Textos organizados por história e por página")
    run.italic = True
    run.font.name = "Georgia"
    run.font.size = Pt(14)
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_before = Pt(240)
    author_run = author.add_run("Miguel Pinto")
    author_run.bold = True
    author_run.font.name = "Georgia"
    author_run.font.size = Pt(13)

    document.add_page_break()
    document.add_heading("Índice", level=1)
    for index, (_, book) in enumerate(books, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        number = paragraph.add_run(f"{index:02d}  ")
        number.bold = True
        number.font.color.rgb = RGBColor(181, 71, 51)
        paragraph.add_run(book["title"])

    text_page_count = 0
    for _, book in books:
        for page_index, page in enumerate(book["pages"], start=1):
            if page["type"] not in {"text", "ending"}:
                continue
            document.add_page_break()
            heading = document.add_paragraph(style="Heading 1")
            heading.paragraph_format.space_after = Pt(6)
            heading.add_run(book["title"])
            label = document.add_paragraph(style="Heading 2")
            label.paragraph_format.space_after = Pt(22)
            label_run = label.add_run(f"Página {page_index}")
            label_run.font.color.rgb = RGBColor(181, 71, 51)

            for paragraph_text in page.get("text", []):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(10)
                paragraph.paragraph_format.line_spacing = 1.25
                paragraph.add_run(paragraph_text)
            if page.get("end"):
                ending = document.add_paragraph()
                ending.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ending.paragraph_format.space_before = Pt(18)
                ending_run = ending.add_run(page["end"])
                ending_run.italic = True
                ending_run.font.color.rgb = RGBColor(181, 71, 51)
            text_page_count += 1

    core_properties = document.core_properties
    core_properties.title = "As Histórias do Cão Joaquim Textos por Página"
    core_properties.author = "Miguel Pinto"
    core_properties.subject = "Textos completos organizados para revisão e impressão"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return text_page_count


def copy_site_and_pdfs(package_dir: Path) -> None:
    pdf_dir = package_dir / "06_PDFS"
    pdf_dir.mkdir(parents=True, exist_ok=True)
    for pdf in sorted(PRINT_DIR.glob("*.pdf")):
        shutil.copy2(pdf, pdf_dir / pdf.name)

    site_dir = package_dir / "07_SITE_COMPLETO"
    site_dir.mkdir(parents=True, exist_ok=True)
    for filename in SITE_ROOT_FILES:
        source = ROOT / filename
        if source.exists():
            shutil.copy2(source, site_dir / filename)
    for directory in SITE_DIRECTORIES:
        source = ROOT / directory
        if source.exists():
            shutil.copytree(source, site_dir / directory, dirs_exist_ok=True)


def write_guide(path: Path, page_count: int, text_count: int, illustration_count: int) -> None:
    content = f"""# Pacote de impressão A4

Este pacote foi preparado para imprimir **As Histórias do Cão Joaquim** em formato A4 com capa dura.

## Conteúdo

- `01_CAPA_DURA`: capa, contracapa, lombadas e capas abertas com sangria;
- `02_PAGINAS_A4_300DPI`: {page_count} páginas completas em JPG, 2480 x 3508 px, 300 dpi e sRGB;
- `03_ILUSTRACOES_A4_300DPI`: {illustration_count} capas e ilustrações separadas em A4, 300 dpi;
- `04_TEXTOS_IMAGEM_A4_300DPI`: {text_count} páginas de texto em PNG, 2480 x 3508 px e 300 dpi;
- `05_TEXTOS_WORD`: textos completos separados por história e página;
- `06_PDFS`: todos os PDFs atuais;
- `07_SITE_COMPLETO`: cópia completa do site, incluindo a V1.

## Interior recomendado

Para um serviço de fotolivros, usa as páginas `002` a `147` da pasta `02_PAGINAS_A4_300DPI` como interior. A página `001` e a página `148` ficam como referência visual, porque a capa dura deve ser enviada separadamente.

As páginas já estão na ordem correta. Não é necessário reorganizá-las nem aplicar outra moldura.

## Capa dura

As capas abertas incluem 3 mm de sangria exterior e estão disponíveis com lombadas de 10, 15, 20, 25 e 30 mm. A largura correta depende do número de folhas, da espessura do papel e do sistema de encadernação da gráfica.

Antes de encomendar, pede ao fornecedor a largura exata da lombada e o modelo da capa. Escolhe o ficheiro correspondente apenas se a medida for igual. Se for diferente, a lombada deve ser ajustada ao modelo da gráfica.

## Definições recomendadas

- formato final: A4 vertical, 210 x 297 mm;
- resolução: 300 dpi;
- espaço de cor: sRGB;
- impressão: frente e verso;
- interior: {page_count - 2} páginas, excluindo a capa e contracapa de referência;
- não ativar correção automática de cor nem melhoria automática das fotografias;
- não cortar nem ampliar automaticamente as páginas;
- confirmar no preview que nenhuma margem ou texto entra na zona de corte.

Se a gráfica pedir CMYK, PDF/X, outra sangria, margens de segurança ou uma lombada diferente, usa o modelo fornecido por ela antes de imprimir.
"""
    path.write_text(content, encoding="utf-8")


def zip_package(source_dir: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".zip.tmp")
    temporary.unlink(missing_ok=True)
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9, allowZip64=True) as archive:
        for path in sorted(item for item in source_dir.rglob("*") if item.is_file()):
            archive.write(path, path.relative_to(source_dir).as_posix())
    temporary.replace(output_path)


def build_print_package() -> Path:
    if not COLLECTION_PDF.exists():
        raise FileNotFoundError(f"PDF da coletanea em falta: {COLLECTION_PDF}")
    if not COVER_ART.exists():
        raise FileNotFoundError(f"Imagem de capa em falta: {COVER_ART}")

    catalog = load_json(BOOKS_DIR / "books.json")
    books: list[tuple[Path, dict]] = []
    for item in catalog["books"]:
        if item.get("status") != "available":
            continue
        book_dir = BOOKS_DIR / item["id"]
        books.append((book_dir, load_json(book_dir / "book.json")))

    if PACKAGE_DIR.exists():
        shutil.rmtree(PACKAGE_DIR)
    PACKAGE_DIR.mkdir(parents=True)

    cover_dir = PACKAGE_DIR / "01_CAPA_DURA"
    pages_dir = PACKAGE_DIR / "02_PAGINAS_A4_300DPI"
    illustrations_dir = PACKAGE_DIR / "03_ILUSTRACOES_A4_300DPI"
    text_images_dir = PACKAGE_DIR / "04_TEXTOS_IMAGEM_A4_300DPI"
    documents_dir = PACKAGE_DIR / "05_TEXTOS_WORD"

    front = make_front_cover(A4_PORTRAIT)
    back = make_back_cover(A4_PORTRAIT, len(books))
    save_jpeg(front, cover_dir / "capa-frente-A4-300dpi.jpg")
    save_jpeg(back, cover_dir / "contracapa-A4-300dpi.jpg")
    save_png(front, cover_dir / "capa-frente-A4-300dpi.png")
    save_png(back, cover_dir / "contracapa-A4-300dpi.png")

    for spine_mm in SPINE_WIDTHS_MM:
        spine_px = round(spine_mm / 25.4 * 300)
        spine = make_spine(spine_px, A4_PORTRAIT[1])
        save_png(spine, cover_dir / f"lombada-{spine_mm:02d}mm-300dpi.png")
        wrap = make_wrap(front, back, spine_mm)
        save_jpeg(wrap, cover_dir / f"capa-aberta-A4-lombada-{spine_mm:02d}mm-sangria-3mm.jpg")

    page_outputs = render_collection_pages(pages_dir)
    illustration_count = generate_illustrations(books, illustrations_dir)
    text_count = generate_text_images(books, text_images_dir)
    docx_path = documents_dir / "as-historias-do-cao-joaquim-textos-por-pagina.docx"
    document_text_count = create_text_document(books, docx_path)
    if document_text_count != text_count:
        raise ValueError("O numero de paginas de texto do Word nao corresponde as imagens.")

    copy_site_and_pdfs(PACKAGE_DIR)
    write_guide(PACKAGE_DIR / "00_LEIA_ME_IMPRESSAO_A4.md", len(page_outputs), text_count, illustration_count)

    manifest = {
        "title": "As Histórias do Cão Joaquim",
        "author": "Miguel Pinto",
        "format": "A4",
        "dpi": 300,
        "color_space": "sRGB",
        "complete_pages": len(page_outputs),
        "interior_pages_recommended": len(page_outputs) - 2,
        "text_pages": text_count,
        "illustrations": illustration_count,
        "stories": len(books),
        "spine_widths_mm": list(SPINE_WIDTHS_MM),
    }
    (PACKAGE_DIR / "manifesto.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    zip_package(PACKAGE_DIR, PACKAGE_ZIP)
    return PACKAGE_ZIP


def main() -> None:
    output = build_print_package()
    print(output.relative_to(ROOT))


if __name__ == "__main__":
    main()
